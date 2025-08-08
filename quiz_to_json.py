import os
import json
import tkinter as tk
from tkinter import filedialog, messagebox
from pathlib import Path

# 解析Excel
import openpyxl

# 解析Word
from docx import Document

# 解析PDF
import fitz  # PyMuPDF

# -------------------------------
# 核心解析函数，返回题目列表，每题为字典：
# { "question": str, "options": {"A": "...", "B": "...", ...}, "answer": str }
# -------------------------------

def parse_xlsx(file_path):
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    questions = []
    for row in ws.iter_rows(min_row=2, values_only=True):  # 假设第一行为表头
        # 列顺序：题目 / A / B / C / D / 答案
        question = str(row[0]).strip() if row[0] else ""
        options = {}
        for idx, opt_label in enumerate(['A', 'B', 'C', 'D']):
            cell_val = row[idx + 1]
            if cell_val is not None:
                options[opt_label] = str(cell_val).strip()
        answer = str(row[5]).strip() if len(row) > 5 and row[5] else ""

        # 非选择题处理：题干有，options空，answer空
        if not options:
            options = {}
            answer = ""

        questions.append({
            "question": question,
            "options": options,
            "answer": answer
        })
    return questions

def parse_docx(file_path):
    doc = Document(file_path)
    questions = []

    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip() != ""]
    # Word格式：
    # 题干行 -> 4个选项行（A/B/C/D）-> 答案行
    i = 0
    while i < len(paragraphs):
        question = paragraphs[i]
        i += 1

        options = {}
        # 读取选项，最多4个，以A、B、C、D开头
        for opt_label in ['A', 'B', 'C', 'D']:
            if i < len(paragraphs) and paragraphs[i].startswith(opt_label):
                # 去除开头的 "A. " 或 "A "等
                opt_text = paragraphs[i][len(opt_label):].lstrip('. ').strip()
                options[opt_label] = opt_text
                i += 1
            else:
                break

        # 读取答案行（可无）
        answer = ""
        if i < len(paragraphs) and paragraphs[i].lower().startswith("答案"):
            ans_line = paragraphs[i]
            # 答案格式假设是“答案: A”或“答案 A”
            parts = ans_line.split(":") if ":" in ans_line else ans_line.split()
            if len(parts) > 1:
                answer = parts[1].strip()
            i += 1

        # 非选择题处理
        if not options:
            options = {}
            answer = ""

        questions.append({
            "question": question,
            "options": options,
            "answer": answer
        })
    return questions

def parse_pdf(file_path):
    # PDF解析比较复杂，依赖排版，这里简化处理：
    # 按行提取文本，使用和Word相似的规则
    doc = fitz.open(file_path)
    texts = []
    for page in doc:
        texts.extend(page.get_text("text").split('\n'))
    paragraphs = [t.strip() for t in texts if t.strip() != ""]

    questions = []
    i = 0
    while i < len(paragraphs):
        question = paragraphs[i]
        i += 1

        options = {}
        for opt_label in ['A', 'B', 'C', 'D']:
            if i < len(paragraphs) and paragraphs[i].startswith(opt_label):
                opt_text = paragraphs[i][len(opt_label):].lstrip('. ').strip()
                options[opt_label] = opt_text
                i += 1
            else:
                break

        answer = ""
        if i < len(paragraphs) and paragraphs[i].lower().startswith("答案"):
            ans_line = paragraphs[i]
            parts = ans_line.split(":") if ":" in ans_line else ans_line.split()
            if len(parts) > 1:
                answer = parts[1].strip()
            i += 1

        if not options:
            options = {}
            answer = ""

        questions.append({
            "question": question,
            "options": options,
            "answer": answer
        })
    return questions

# -------------------------------
# 保存 JSON 文件
# -------------------------------
def save_json(questions, output_path):
    with open(output_path, 'w', encoding='utf-8') as f:
        json.dump(questions, f, ensure_ascii=False, indent=2)

# -------------------------------
# GUI 相关
# -------------------------------
class App:
    def __init__(self, root):
        self.root = root
        root.title("题库转 JSON 工具 (极简版)")
        root.geometry("480x150")
        root.resizable(False, False)

        self.file_path = ""

        # 选择文件按钮 + 显示路径标签
        self.btn_select = tk.Button(root, text="选择文件", width=15, command=self.select_file)
        self.btn_select.pack(pady=10)

        self.lbl_path = tk.Label(root, text="未选择文件", fg="blue")
        self.lbl_path.pack()

        # 生成题库按钮
        self.btn_generate = tk.Button(root, text="生成题库", width=15, command=self.generate_json, state='disabled')
        self.btn_generate.pack(pady=15)

    def select_file(self):
        f = filedialog.askopenfilename(
            filetypes=[("题库文件", "*.xlsx *.docx *.pdf")],
            title="请选择 Excel/Word/PDF 题库文件"
        )
        if f:
            self.file_path = f
            self.lbl_path.config(text=f)
            self.btn_generate.config(state='normal')

    def generate_json(self):
        if not self.file_path:
            messagebox.showwarning("警告", "请先选择文件！")
            return

        ext = Path(self.file_path).suffix.lower()
        try:
            if ext == '.xlsx':
                questions = parse_xlsx(self.file_path)
            elif ext == '.docx':
                questions = parse_docx(self.file_path)
            elif ext == '.pdf':
                questions = parse_pdf(self.file_path)
            else:
                messagebox.showerror("错误", "不支持的文件格式！")
                return
        except Exception as e:
            messagebox.showerror("错误", f"解析文件失败:\n{str(e)}")
            return

        # 保存 JSON
        json_path = str(Path(self.file_path).with_suffix('.json'))
        try:
            save_json(questions, json_path)
        except Exception as e:
            messagebox.showerror("错误", f"保存 JSON 失败:\n{str(e)}")
            return

        messagebox.showinfo("成功", f"已生成 JSON 文件:\n{json_path}")

# -------------------------------
# 主程序启动
# -------------------------------
def main():
    root = tk.Tk()
    app = App(root)
    root.mainloop()

if __name__ == "__main__":
    main()
